import docx
import docx.shared
from docx.shared import Inches


def main():
    doc = docx.Document()
    doc.add_heading("Бесполезная информация", level=1)

    p1 = doc.add_paragraph("Ленивцы могут задерживать дыхание под водой до 40 минут")
    style = p1.runs[0]
    style.font.size = docx.shared.Pt(15)
    style.font.name = "Times New Roman"
    style.bold = True

    p2 = doc.add_paragraph("Белый медведь это единственное животное, которое не видно в тепловизор.")
    style = p2.runs[0]
    style.font.size = docx.shared.Pt(20)
    style.font.name = "Cambria"
    style.italic = True

    doc.add_picture("./white_bear.jpg", width=Inches(4), height=Inches(3))

    p3 = doc.add_paragraph("А хотя нет, все же видно")
    style = p3.runs[0]
    style.font.size = docx.shared.Pt(12)
    style.font.name = "Arial Black"
    style.underline = True

    doc.save("useless facts.docx")


if __name__ == "__main__":
    main()
